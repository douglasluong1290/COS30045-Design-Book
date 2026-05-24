"""
One-shot xlsx → JSON converter for the storytelling page.
Reads thanh_analysis/datasets/transformed_datasets.xlsx and writes
the JSON files D3 charts consume into public/data/.

Run from repo root:
    python3 scripts/build_data.py
"""

import json
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent.parent
SRC  = REPO / "thanh_analysis" / "datasets" / "transformed_datasets.xlsx"
OUT  = REPO / "public" / "data"
OUT.mkdir(parents=True, exist_ok=True)


def admission_category(year: int) -> str:
    """Mirrors the DAX SWITCH in charts_features.txt."""
    if year == 2011:
        return "Hospitalised injuries"
    if 2012 <= year <= 2016:
        return "Change in admissions 2012"
    return "Change in admissions 2017"


def to_int(v):
    if v in (None, "", "n.p.", "np"):
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def to_float(v):
    if v in (None, "", "n.p.", "np"):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def main() -> None:
    wb = load_workbook(SRC, data_only=True)

    # ---------------- chart 1 — national trend ----------------
    pub = wb["publication"]
    by_year: dict[int, int] = defaultdict(int)
    for row in pub.iter_rows(min_row=2, values_only=True):
        year = to_int(row[0])
        hosp = to_int(row[12])  # Hospitalisations column
        if year is None or hosp is None:
            continue
        by_year[year] += hosp
    national_trend = [
        {"year": y, "hospitalisations": by_year[y], "admission": admission_category(y)}
        for y in sorted(by_year)
    ]
    (OUT / "national_trend.json").write_text(json.dumps(national_trend, indent=2))
    print(f"national_trend.json — {len(national_trend)} rows")

    # ---------------- chart 2 — state totals + drill data ----------------
    state_summary = wb["state_summary"]
    state_year: dict[tuple[str, int], dict] = defaultdict(
        lambda: {"cases": 0, "bed_days": 0}
    )
    for row in state_summary.iter_rows(min_row=2, values_only=True):
        year = to_int(row[0])
        state = row[1]
        cases = to_int(row[2]) or 0
        beds = to_int(row[3]) or 0
        if year is None or not state:
            continue
        state_year[(state, year)]["cases"] += cases
        state_year[(state, year)]["bed_days"] += beds

    state_trend = [
        {"state": st, "year": y, "cases": v["cases"], "bed_days": v["bed_days"]}
        for (st, y), v in sorted(state_year.items())
    ]
    (OUT / "state_trend.json").write_text(json.dumps(state_trend, indent=2))
    print(f"state_trend.json — {len(state_trend)} rows")

    state_totals: dict[str, int] = defaultdict(int)
    for r in state_trend:
        state_totals[r["state"]] += r["cases"]
    state_totals_out = [
        {"state": s, "cases": c}
        for s, c in sorted(state_totals.items(), key=lambda kv: -kv[1])
    ]
    (OUT / "state_totals.json").write_text(json.dumps(state_totals_out, indent=2))
    print(f"state_totals.json — {len(state_totals_out)} rows")

    # population by state per year — for the drill bar
    pop_sheet = wb["population"]
    state_alias = {
        "New South Wales": "NSW",
        "Victoria": "VIC",
        "Queensland": "QLD",
        "South Australia": "SA",
        "Western Australia": "WA",
        "Tasmania": "TAS",
        "Northern Territory": "NT",
        "Australian Capital Territory": "ACT",
    }
    population: list[dict] = []
    for row in pop_sheet.iter_rows(min_row=2, values_only=True):
        year = to_int(row[0])
        region = row[1]
        pop = to_float(row[2])
        if year is None or pop is None or not region:
            continue
        code = state_alias.get(region, region)
        population.append({"state": code, "year": year, "population": pop})
    (OUT / "population.json").write_text(json.dumps(population, indent=2))
    print(f"population.json — {len(population)} rows")

    # ---------------- chart 3 — age × sex ----------------
    age_sex: dict[tuple[str, str], dict] = defaultdict(
        lambda: {"cases": 0, "bed_days": 0}
    )
    for row in pub.iter_rows(min_row=2, values_only=True):
        age = row[4]
        sex = row[5]
        cases = to_int(row[8]) or 0
        beds = to_int(row[9]) or 0
        if not age or not sex:
            continue
        # drop the suppressed / missing age band per spec
        if str(age).strip().lower() in {"missing", "not stated", "unknown"}:
            continue
        age_sex[(str(age), str(sex))]["cases"] += cases
        age_sex[(str(age), str(sex))]["bed_days"] += beds

    age_sex_out = [
        {"age_band": a, "sex": s, "cases": v["cases"], "bed_days": v["bed_days"]}
        for (a, s), v in age_sex.items()
    ]
    (OUT / "age_sex.json").write_text(json.dumps(age_sex_out, indent=2))
    print(f"age_sex.json — {len(age_sex_out)} rows")

    # ---------------- chart 4 — road user ----------------
    road_user: dict[str, dict] = defaultdict(lambda: {"cases": 0, "bed_days": 0})
    for row in pub.iter_rows(min_row=2, values_only=True):
        ru = row[6]
        cases = to_int(row[8]) or 0
        beds = to_int(row[9]) or 0
        if not ru:
            continue
        if str(ru).strip().lower() == "not applicable":
            # baked-in filter mirrors the page-level filter in the final pbix
            continue
        road_user[str(ru)]["cases"] += cases
        road_user[str(ru)]["bed_days"] += beds

    road_user_out = [
        {"road_user": k, "cases": v["cases"], "bed_days": v["bed_days"]}
        for k, v in sorted(road_user.items(), key=lambda kv: -kv[1]["cases"])
    ]
    (OUT / "road_user.json").write_text(json.dumps(road_user_out, indent=2))
    print(f"road_user.json — {len(road_user_out)} rows")


if __name__ == "__main__":
    main()
